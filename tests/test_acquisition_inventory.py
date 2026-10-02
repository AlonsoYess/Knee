"""Tests for the cohort-level baseline acquisition inventory."""

import csv
import io
import json
import tarfile
import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from pydicom.dataset import FileDataset, FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian, generate_uid

from knee.acquisition_inventory import build_inventory, write_inventory


def dicom_bytes(seed: int) -> bytes:
    pixels = np.arange(24, dtype=np.uint16).reshape(4, 6) + seed
    meta = FileMetaDataset()
    meta.MediaStorageSOPClassUID = generate_uid()
    meta.MediaStorageSOPInstanceUID = generate_uid()
    meta.TransferSyntaxUID = ExplicitVRLittleEndian
    dataset = FileDataset(None, {}, file_meta=meta, preamble=b"\0" * 128)
    dataset.SOPClassUID = meta.MediaStorageSOPClassUID
    dataset.SOPInstanceUID = meta.MediaStorageSOPInstanceUID
    dataset.Rows, dataset.Columns = pixels.shape
    dataset.SamplesPerPixel = 1
    dataset.PhotometricInterpretation = "MONOCHROME2"
    dataset.PixelRepresentation = 0
    dataset.BitsAllocated = 16
    dataset.BitsStored = 16
    dataset.HighBit = 15
    dataset.NumberOfFrames = 1
    dataset.Modality = "CR"
    dataset.PixelData = pixels.tobytes()
    stream = io.BytesIO()
    dataset.save_as(stream, enforce_file_format=True)
    return stream.getvalue()


def write_archive(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    mode = "w:gz" if path.name.endswith(".gz") else "w"
    with tarfile.open(path, mode) as archive:
        member = tarfile.TarInfo("image.dcm")
        member.size = len(payload)
        archive.addfile(member, io.BytesIO(payload))


class AcquisitionInventoryTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.manifest = self.root / "manifest.xlsx"
        self.pilot = self.root / "pilot_v1"
        self.cohort = self.root / "cohorte_v00"
        self.pilot.mkdir()
        self.cohort.mkdir()
        self._write_manifest()

    def _write_manifest(self) -> None:
        workbook = Workbook()
        unique = workbook.active
        unique.title = "Imagenes unicas"
        unique.append(
            [
                "SRC_SUBJECT_ID",
                "BARCODE_BASE",
                "IMAGE_FILE",
                "IMAGE_THUMBNAIL_FILE",
                "RODILLAS_ASOCIADAS",
                "LATERALIDADES",
            ]
        )
        unique.append(["P1", "10000001", "path/image_10000001.tar.gz", "thumb/1", 2, "D+I"])
        unique.append(["P2", "10000002", "path/image_10000002.tar.gz", "thumb/2", 1, "D"])
        unique.append(["P3", "10000003", "path/image_10000003.tar.gz", "thumb/3", 1, "I"])

        knees = workbook.create_sheet("Manifiesto rodillas")
        knees.append(["SRC_SUBJECT_ID", "SIDE", "BARCODE_BASE", "IMAGE_FILE", "KL_48"])
        knees.append(["P1", 1, "10000001", "path/image_10000001.tar.gz", 3])
        knees.append(["P1", 2, "10000001", "path/image_10000001.tar.gz", 2])
        knees.append(["P2", 1, "10000002", "path/image_10000002.tar.gz", 4])
        knees.append(["P3", 2, "10000003", "path/image_10000003.tar.gz", 3])
        workbook.save(self.manifest)

    def test_inventory_classifies_available_duplicate_and_pending(self):
        payload = dicom_bytes(1)
        write_archive(self.pilot / "image_10000001.tar.gz", payload)
        write_archive(self.cohort / "copy" / "image_10000001.tar", payload)
        write_archive(self.cohort / "image_10000002.tar.gz", dicom_bytes(2))

        result = build_inventory(
            self.manifest,
            [self.pilot, self.cohort],
            expected_unique_acquisitions=3,
            download_batch_size=2,
            workers=2,
        )
        summary = result["public_summary"]

        self.assertEqual(summary["status"], "attention_required")
        self.assertEqual(
            summary["local_status_counts"],
            {
                "DISPONIBLE": 1,
                "DUPLICADO_EQUIVALENTE": 1,
                "PENDIENTE_DESCARGA": 1,
            },
        )
        self.assertEqual(summary["pending_download_or_replacement"], 1)
        self.assertEqual(result["download_queue"][0]["download_batch"], 1)
        self.assertFalse(summary["training_executed"])
        self.assertFalse(summary["reserved_test_opened"])

    def test_clean_partial_inventory_is_ready_for_selective_download(self):
        write_archive(self.pilot / "image_10000001.tar.gz", dicom_bytes(1))
        result = build_inventory(
            self.manifest,
            [self.pilot, self.cohort],
            expected_unique_acquisitions=3,
            download_batch_size=1,
        )

        self.assertEqual(result["public_summary"]["status"], "ready_for_selective_download")
        self.assertEqual(result["public_summary"]["download_batches"], 2)
        self.assertEqual(
            [row["download_batch"] for row in result["download_queue"]], [1, 2]
        )

    def test_output_separates_public_counts_from_private_identifiers(self):
        result = build_inventory(
            self.manifest,
            [self.pilot, self.cohort],
            expected_unique_acquisitions=3,
        )
        output = self.root / "output"
        write_inventory(result, output)
        public_text = (output / "resumen_inventario_publico.json").read_text(
            encoding="utf-8"
        )
        with (output / "inventario_adquisiciones_privado.csv").open(
            encoding="utf-8-sig", newline=""
        ) as stream:
            rows = list(csv.DictReader(stream, delimiter=";"))

        self.assertNotIn("P1", public_text)
        self.assertNotIn("image_10000001", public_text)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]["src_subject_id"], "P1")
        self.assertTrue((output / "plan_descarga_congelado_privado.csv").is_file())

    def test_frozen_plan_keeps_batch_numbers_after_a_completed_lot(self):
        write_archive(self.pilot / "image_10000001.tar.gz", dicom_bytes(1))
        initial = build_inventory(
            self.manifest,
            [self.pilot, self.cohort],
            expected_unique_acquisitions=3,
            download_batch_size=1,
        )
        self.assertEqual(
            [row["download_batch"] for row in initial["download_queue"]], [1, 2]
        )

        write_archive(self.cohort / "image_10000002.tar.gz", dicom_bytes(2))
        updated = build_inventory(
            self.manifest,
            [self.pilot, self.cohort],
            expected_unique_acquisitions=3,
            download_batch_size=1,
            batch_plan=initial["download_plan"],
        )

        self.assertEqual(len(updated["download_queue"]), 1)
        self.assertEqual(updated["download_queue"][0]["download_batch"], 2)
        self.assertEqual(updated["public_summary"]["remaining_batch_numbers"], [2])
        self.assertEqual(len(updated["download_plan"]), 2)

    def test_manifest_mismatch_stops_before_file_audit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Imagenes unicas"
        sheet.append(list((
            "SRC_SUBJECT_ID",
            "BARCODE_BASE",
            "IMAGE_FILE",
            "IMAGE_THUMBNAIL_FILE",
            "RODILLAS_ASOCIADAS",
            "LATERALIDADES",
        )))
        sheet.append(["P1", "10000001", "wrong_name.tar.gz", "thumb", 1, "D"])
        knees = workbook.create_sheet("Rodillas")
        knees.append(["SRC_SUBJECT_ID", "SIDE", "BARCODE_BASE", "IMAGE_FILE"])
        knees.append(["P1", 1, "10000001", "wrong_name.tar.gz"])
        workbook.save(self.manifest)

        with self.assertRaisesRegex(ValueError, "image_file_barcode_mismatch"):
            build_inventory(
                self.manifest,
                [self.pilot, self.cohort],
                expected_unique_acquisitions=1,
            )


if __name__ == "__main__":
    unittest.main()
