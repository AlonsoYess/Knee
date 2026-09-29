"""Build the private baseline-acquisition inventory for selective OAI download."""

from __future__ import annotations

import argparse
import csv
import json
import os
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from knee.dicom_audit import (
    DETAIL_FIELDS,
    _archive_paths,
    _audit_archive,
    _selection_rank,
    canonical_package_key,
    sha256_file,
)


UNIQUE_REQUIRED = (
    "SRC_SUBJECT_ID",
    "BARCODE_BASE",
    "IMAGE_FILE",
    "IMAGE_THUMBNAIL_FILE",
    "RODILLAS_ASOCIADAS",
    "LATERALIDADES",
)
KNEE_REQUIRED = ("SRC_SUBJECT_ID", "SIDE", "BARCODE_BASE", "IMAGE_FILE")
INVENTORY_FIELDS = (
    "inventory_index",
    "acquisition_key",
    "src_subject_id",
    "barcode_base",
    "image_file",
    "image_thumbnail_file",
    "associated_knees",
    "lateralities",
    "local_status",
    "action_required",
    "download_batch",
    "archive_count",
    "readable_archive_count",
    "distinct_pixel_contents",
    "canonical_source",
    "canonical_package_relative_path",
    "package_bytes",
    "package_sha256",
    "dicom_sha256",
    "pixel_sha256",
)
DOWNLOAD_FIELDS = (
    "download_batch",
    "inventory_index",
    "acquisition_key",
    "image_file",
    "local_status",
    "action_required",
)
DOWNLOAD_PLAN_FIELDS = (
    "download_batch",
    "inventory_index",
    "acquisition_key",
    "image_file",
)
UNEXPECTED_FIELDS = (
    "source",
    "manifest_key",
    "package_relative_path",
    "package_bytes",
    "package_sha256",
    "dicom_valid",
    "error_code",
)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(sheet: Any, required: Iterable[str]) -> list[dict[str, str]]:
    values = sheet.iter_rows(values_only=True)
    try:
        header = [_cell_text(value) for value in next(values)]
    except StopIteration as exc:
        raise ValueError(f"Empty worksheet: {sheet.title}") from exc
    if len(header) != len(set(header)):
        raise ValueError(f"Duplicated header in worksheet: {sheet.title}")
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"Worksheet {sheet.title} lacks columns: {', '.join(missing)}")
    positions = {name: header.index(name) for name in required}
    rows: list[dict[str, str]] = []
    for raw in values:
        row = {name: _cell_text(raw[index]) for name, index in positions.items()}
        if any(row.values()):
            rows.append(row)
    return rows


def _find_sheet(workbook: Any, required: Iterable[str]) -> Any:
    required_set = set(required)
    matches = []
    for sheet in workbook.worksheets:
        header = {
            _cell_text(cell.value)
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        }
        if required_set.issubset(header):
            matches.append(sheet)
    if len(matches) != 1:
        raise ValueError(
            "Expected exactly one worksheet with columns "
            f"{sorted(required_set)}; found {len(matches)}."
        )
    return matches[0]


def read_manifest_contract(
    manifest_xlsx: Path, expected_unique_acquisitions: int
) -> tuple[list[dict[str, str]], dict[str, Any]]:
    """Read only the identity columns needed for the image inventory."""
    workbook = load_workbook(manifest_xlsx, read_only=True, data_only=True)
    try:
        unique_rows = _read_rows(_find_sheet(workbook, UNIQUE_REQUIRED), UNIQUE_REQUIRED)
        knee_rows = _read_rows(_find_sheet(workbook, KNEE_REQUIRED), KNEE_REQUIRED)
    finally:
        workbook.close()

    errors: list[str] = []
    for column in UNIQUE_REQUIRED:
        if any(not row[column] for row in unique_rows):
            errors.append(f"blank_{column.lower()}")

    acquisition_keys = [canonical_package_key(row["IMAGE_FILE"]) for row in unique_rows]
    identities = [
        (row["SRC_SUBJECT_ID"], row["BARCODE_BASE"], row["IMAGE_FILE"])
        for row in unique_rows
    ]
    if len(unique_rows) != expected_unique_acquisitions:
        errors.append("unexpected_unique_acquisition_count")
    if len(set(acquisition_keys)) != len(acquisition_keys):
        errors.append("duplicate_acquisition_key")
    if len(set(identities)) != len(identities):
        errors.append("duplicate_manifest_identity")
    if any(
        not key.endswith(str(row["BARCODE_BASE"])[-8:])
        for key, row in zip(acquisition_keys, unique_rows)
    ):
        errors.append("image_file_barcode_mismatch")

    knee_groups: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for row in knee_rows:
        identity = (row["SRC_SUBJECT_ID"], row["BARCODE_BASE"], row["IMAGE_FILE"])
        knee_groups[identity].append(row["SIDE"])
    if set(knee_groups) != set(identities):
        errors.append("knee_manifest_identity_mismatch")
    for row, identity in zip(unique_rows, identities):
        try:
            expected_knees = int(row["RODILLAS_ASOCIADAS"])
        except ValueError:
            errors.append("invalid_associated_knee_count")
            continue
        if expected_knees != len(knee_groups.get(identity, [])):
            errors.append("associated_knee_count_mismatch")

    if errors:
        raise ValueError("Manifest contract failed: " + ", ".join(sorted(set(errors))))

    records = []
    for index, (row, key) in enumerate(zip(unique_rows, acquisition_keys), start=1):
        records.append(
            {
                "inventory_index": str(index),
                "acquisition_key": key,
                "src_subject_id": row["SRC_SUBJECT_ID"],
                "barcode_base": row["BARCODE_BASE"],
                "image_file": row["IMAGE_FILE"],
                "image_thumbnail_file": row["IMAGE_THUMBNAIL_FILE"],
                "associated_knees": row["RODILLAS_ASOCIADAS"],
                "lateralities": row["LATERALIDADES"],
            }
        )

    contract = {
        "status": "ok",
        "unique_acquisitions": len(records),
        "participants": len({row["src_subject_id"] for row in records}),
        "knee_rows": len(knee_rows),
        "acquisitions_with_one_eligible_knee": sum(
            row["associated_knees"] == "1" for row in records
        ),
        "acquisitions_with_two_eligible_knees": sum(
            row["associated_knees"] == "2" for row in records
        ),
    }
    return records, contract


def _audit_sources(source_dirs: list[Path], workers: int) -> list[dict[str, Any]]:
    tasks: list[tuple[str, Path, Path]] = []
    for source_dir in source_dirs:
        if not source_dir.is_dir():
            raise FileNotFoundError(f"DICOM source directory not found: {source_dir}")
        for path in _archive_paths(source_dir):
            tasks.append((source_dir.name, source_dir, path))

    def audit(task: tuple[str, Path, Path]) -> dict[str, Any]:
        source_name, root, path = task
        record = _audit_archive(path, root)
        record["source"] = source_name
        return record

    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        return list(executor.map(audit, tasks))


def _classify(candidates: list[dict[str, Any]]) -> tuple[str, str, dict[str, Any] | None]:
    if not candidates:
        return "PENDIENTE_DESCARGA", "DESCARGAR", None
    valid = [record for record in candidates if record["dicom_valid"]]
    identities = {str(record["pixel_sha256"]) for record in valid}
    if not valid:
        return "ILEGIBLE", "REEMPLAZAR", None
    if len(identities) > 1:
        return "DUPLICADO_CONFLICTIVO", "REVISAR", None
    selected = min(valid, key=_selection_rank)
    if len(candidates) == 1:
        return "DISPONIBLE", "NINGUNA", selected
    if len(valid) == len(candidates):
        return "DUPLICADO_EQUIVALENTE", "DEPURAR_COPIAS", selected
    return "DUPLICADO_PARCIALMENTE_ILEGIBLE", "REVISAR", selected


def load_batch_plan(path: Path) -> list[dict[str, str]]:
    with Path(path).open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream, delimiter=";")
        if not reader.fieldnames or not set(DOWNLOAD_PLAN_FIELDS).issubset(
            reader.fieldnames
        ):
            raise ValueError("The frozen download plan has an invalid schema.")
        rows = [
            {name: str(row.get(name, "")).strip() for name in DOWNLOAD_PLAN_FIELDS}
            for row in reader
        ]
    keys = [row["acquisition_key"] for row in rows]
    urls = [row["image_file"] for row in rows]
    if len(keys) != len(set(keys)) or len(urls) != len(set(urls)):
        raise ValueError("The frozen download plan contains duplicate identities.")
    for row in rows:
        if not row["download_batch"].isdigit() or int(row["download_batch"]) < 1:
            raise ValueError("The frozen download plan contains an invalid batch number.")
        if canonical_package_key(row["image_file"]) != row["acquisition_key"]:
            raise ValueError("The frozen plan key does not match its S3 URL.")
    return rows


def build_inventory(
    manifest_xlsx: Path,
    source_dirs: list[Path],
    expected_unique_acquisitions: int = 1916,
    download_batch_size: int = 100,
    workers: int = 4,
    batch_plan: list[dict[str, str]] | None = None,
) -> dict[str, Any]:
    """Create a deterministic private inventory without reading outcome labels."""
    if download_batch_size < 1:
        raise ValueError("download_batch_size must be positive.")
    expected, contract = read_manifest_contract(
        Path(manifest_xlsx), expected_unique_acquisitions
    )
    audited = _audit_sources([Path(path) for path in source_dirs], workers)
    by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in audited:
        by_key[str(record["manifest_key"])].append(record)

    expected_keys = {row["acquisition_key"] for row in expected}
    inventory: list[dict[str, Any]] = []
    download_queue: list[dict[str, Any]] = []
    for row in expected:
        candidates = by_key.get(row["acquisition_key"], [])
        status, action, selected = _classify(candidates)
        valid = [record for record in candidates if record["dicom_valid"]]
        item: dict[str, Any] = {
            **row,
            "local_status": status,
            "action_required": action,
            "download_batch": "",
            "archive_count": len(candidates),
            "readable_archive_count": len(valid),
            "distinct_pixel_contents": len(
                {str(record["pixel_sha256"]) for record in valid}
            ),
            "canonical_source": "",
            "canonical_package_relative_path": "",
            "package_bytes": "",
            "package_sha256": "",
            "dicom_sha256": "",
            "pixel_sha256": "",
        }
        if selected:
            item.update(
                {
                    "canonical_source": selected["source"],
                    "canonical_package_relative_path": selected["package_relative_path"],
                    "package_bytes": selected["package_bytes"],
                    "package_sha256": selected["package_sha256"],
                    "dicom_sha256": selected["dicom_sha256"],
                    "pixel_sha256": selected["pixel_sha256"],
                }
            )
        if action in {"DESCARGAR", "REEMPLAZAR"}:
            download_queue.append(item)
        inventory.append(item)

    expected_by_key = {row["acquisition_key"]: row for row in expected}
    plan_by_key: dict[str, dict[str, str]] = {}
    if batch_plan:
        for row in batch_plan:
            key = row["acquisition_key"]
            if key not in expected_by_key:
                raise ValueError("The frozen download plan contains an unexpected key.")
            if row["image_file"] != expected_by_key[key]["image_file"]:
                raise ValueError("The frozen download plan no longer matches the manifest.")
            plan_by_key[key] = dict(row)

    new_plan_rows: list[dict[str, str]] = []
    unassigned: list[dict[str, Any]] = []
    for item in download_queue:
        planned = plan_by_key.get(str(item["acquisition_key"]))
        if planned:
            item["download_batch"] = int(planned["download_batch"])
        else:
            unassigned.append(item)
    next_batch = max(
        (int(row["download_batch"]) for row in plan_by_key.values()), default=0
    )
    for position, item in enumerate(unassigned):
        item["download_batch"] = next_batch + 1 + position // download_batch_size
        new_plan_rows.append(
            {
                "download_batch": str(item["download_batch"]),
                "inventory_index": str(item["inventory_index"]),
                "acquisition_key": str(item["acquisition_key"]),
                "image_file": str(item["image_file"]),
            }
        )
    frozen_plan = sorted(
        [*plan_by_key.values(), *new_plan_rows],
        key=lambda row: int(row["inventory_index"]),
    )

    unexpected = [
        record for record in audited if str(record["manifest_key"]) not in expected_keys
    ]
    status_counts = dict(sorted(Counter(row["local_status"] for row in inventory).items()))
    attention = sum(
        count
        for status, count in status_counts.items()
        if status not in {"DISPONIBLE", "PENDIENTE_DESCARGA"}
    )
    public_status = "attention_required" if attention or unexpected else (
        "complete" if not download_queue else "ready_for_selective_download"
    )
    remaining_batches = sorted(
        {int(row["download_batch"]) for row in download_queue}
    )
    public_summary = {
        "status": public_status,
        "manifest_contract": contract,
        "expected_unique_acquisitions": expected_unique_acquisitions,
        "archives_scanned": len(audited),
        "local_status_counts": status_counts,
        "pending_download_or_replacement": len(download_queue),
        "download_batch_size": download_batch_size,
        "download_batches": len(remaining_batches),
        "remaining_batch_numbers": remaining_batches,
        "first_pending_batch": remaining_batches[0] if remaining_batches else None,
        "unexpected_archives": len(unexpected),
        "training_executed": False,
        "reserved_test_opened": False,
    }
    return {
        "public_summary": public_summary,
        "private_metadata": {
            "manifest_sha256": sha256_file(Path(manifest_xlsx)),
            "source_directories": [str(path) for path in source_dirs],
        },
        "inventory": inventory,
        "download_queue": download_queue,
        "download_plan": frozen_plan,
        "unexpected": unexpected,
        "archive_records": audited,
    }


def _write_csv(path: Path, fieldnames: Iterable[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_inventory(result: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "resumen_inventario_publico.json").write_text(
        json.dumps(result["public_summary"], indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (output_dir / "metadatos_inventario_privado.json").write_text(
        json.dumps(result["private_metadata"], indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    _write_csv(
        output_dir / "inventario_adquisiciones_privado.csv",
        INVENTORY_FIELDS,
        result["inventory"],
    )
    _write_csv(
        output_dir / "cola_descarga_selectiva_privada.csv",
        DOWNLOAD_FIELDS,
        result["download_queue"],
    )
    _write_csv(
        output_dir / "plan_descarga_congelado_privado.csv",
        DOWNLOAD_PLAN_FIELDS,
        result["download_plan"],
    )
    _write_csv(
        output_dir / "archivos_inesperados_privado.csv",
        UNEXPECTED_FIELDS,
        result["unexpected"],
    )
    _write_csv(
        output_dir / "auditoria_archivos_privada.csv",
        ("source", *DETAIL_FIELDS),
        result["archive_records"],
    )


def _resolve_under_root(root: Path, value: str, key: str) -> Path:
    relative = Path(value)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError(f"{key} must be relative to KNEE_DATA_ROOT.")
    resolved = (root / relative).resolve()
    if not resolved.is_relative_to(root):
        raise ValueError(f"{key} escapes KNEE_DATA_ROOT.")
    return resolved


def load_inventory_config(path: Path) -> dict[str, Any]:
    root_value = os.environ.get("KNEE_DATA_ROOT")
    if not root_value:
        raise ValueError("Set KNEE_DATA_ROOT to the authorized data directory.")
    root = Path(root_value).expanduser().resolve()
    config = json.loads(Path(path).read_text(encoding="utf-8"))
    for key in ("manifest_xlsx", "output_dir"):
        if not isinstance(config.get(key), str) or not config[key].strip():
            raise ValueError(f"Missing relative path: {key}.")
        config[key] = _resolve_under_root(root, config[key], key)
    source_values = config.get("source_dirs")
    if not isinstance(source_values, list) or not source_values:
        raise ValueError("source_dirs must be a non-empty list of relative paths.")
    config["source_dirs"] = [
        _resolve_under_root(root, value, "source_dirs") for value in source_values
    ]
    config.setdefault("expected_unique_acquisitions", 1916)
    config.setdefault("download_batch_size", 100)
    config.setdefault("workers", 4)
    plan_value = config.get("download_plan_csv")
    if plan_value:
        if not isinstance(plan_value, str):
            raise ValueError("download_plan_csv must be a relative path.")
        config["download_plan_csv"] = _resolve_under_root(
            root, plan_value, "download_plan_csv"
        )
    else:
        config["download_plan_csv"] = (
            config["output_dir"] / "plan_descarga_congelado_privado.csv"
        )
    return config


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", required=True, type=Path)
    args = parser.parse_args()
    config = load_inventory_config(args.config)
    batch_plan = None
    if config["download_plan_csv"].is_file():
        batch_plan = load_batch_plan(config["download_plan_csv"])
    else:
        legacy_queue = config["output_dir"] / "cola_descarga_selectiva_privada.csv"
        if legacy_queue.is_file():
            batch_plan = load_batch_plan(legacy_queue)
    result = build_inventory(
        config["manifest_xlsx"],
        config["source_dirs"],
        int(config["expected_unique_acquisitions"]),
        int(config["download_batch_size"]),
        int(config["workers"]),
        batch_plan,
    )
    write_inventory(result, config["output_dir"])
    print(json.dumps(result["public_summary"], indent=2, ensure_ascii=False))
    if result["public_summary"]["status"] == "attention_required":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
